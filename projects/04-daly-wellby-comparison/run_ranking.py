"""Rank the same charities in life-years per dollar and in WELLBYs per dollar.

Life-saving programs: GiveWell cost per under-5 death, converted to years of
life, then to WELLBYs by a factor this script sweeps.
Cash and psychotherapy: Happier Lives Institute WELLBY totals, used as published.
They avert no deaths in this model, so their life-year count is zero.
"""

from __future__ import annotations

import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
import yaml

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
FIG = HERE / "figures"
GIVEWELL = ROOT / "data" / "raw" / "givewell"

# Simple CEA row for "Cost per under-five death averted", and the first
# geography column. Values left of that column are labels, not countries.
COST_CELLS = {
    "hki_niger": ("hki-vas-cea.xlsx", 19, 9),
    "smc_sokoto": ("malaria-consortium-cea.xlsx", 20, 9),
    "amf_guinea": ("amf-cea.xlsx", 18, 9),
}


def median_cost_per_death(charity_id: str) -> tuple[float, int]:
    from openpyxl import load_workbook

    filename, row, first_col = COST_CELLS[charity_id]
    wb = load_workbook(GIVEWELL / filename, data_only=True, read_only=True)
    values = list(wb["Simple CEA"].iter_rows(min_row=row, max_row=row, values_only=True))[0]
    costs = [float(v) for i, v in enumerate(values, 1) if i >= first_col and isinstance(v, (int, float))]
    return float(pd.Series(costs).median()), len(costs)


def load_conversions(path: Path | None = None) -> dict:
    path = path or HERE / "conversions.yaml"
    return yaml.safe_load(path.read_text())


def years_per_death(spec: dict, charity: dict) -> float:
    base = float(spec["life_years"]["years_per_under5_malaria_death"])
    malaria = float(spec["life_years"]["givewell_u5_malaria_doublings"])
    vitamin = float(spec["life_years"]["givewell_u5_vitamin_a_doublings"])
    if charity["death_weight"] == "vitamin_a":
        return base * vitamin / malaria
    return base


def score(spec: dict, wellbys_per_year: float, cost_basis: str = "best") -> pd.DataFrame:
    rows = []
    for charity in spec["charities"]:
        if charity["kind"] == "life_saving":
            years = years_per_death(spec, charity)
            if cost_basis == "median":
                cost, n_geos = median_cost_per_death(charity["id"])
            else:
                cost, n_geos = float(charity["cost_per_u5_death_usd"]), 1
            life_years_per_1000 = years / cost * 1000
            wellbys_per_1000 = life_years_per_1000 * wellbys_per_year
            cost_out = cost
        else:
            n_geos = None
            years = 0.0
            life_years_per_1000 = 0.0
            wellbys_per_1000 = float(charity["wellbys_per_1000_usd"])
            cost_out = None
        geography = charity["geography"]
        if charity["kind"] == "life_saving" and cost_basis == "median":
            geography = f"Median of {n_geos} geographies"
        rows.append(
            {
                "id": charity["id"],
                "charity": charity["name"],
                "program": charity["program"],
                "geography": geography,
                "kind": charity["kind"],
                "cost_basis": cost_basis if charity["kind"] == "life_saving" else "published",
                "n_geographies": n_geos,
                "cost_per_u5_death_usd": cost_out,
                "years_per_death": years,
                "life_years_per_1000": life_years_per_1000,
                "wellbys_per_1000": wellbys_per_1000,
            }
        )
    out = pd.DataFrame(rows)
    out["rank_life_years"] = out["life_years_per_1000"].rank(ascending=False, method="min").astype(int)
    out["rank_wellby"] = out["wellbys_per_1000"].rank(ascending=False, method="min").astype(int)
    return out.sort_values(["rank_wellby", "charity"])


def sweep(spec: dict) -> pd.DataFrame:
    w = spec["wellby"]
    values = []
    x = float(w["sweep_min"])
    while x <= float(w["sweep_max"]) + 1e-9:
        values.append(round(x, 4))
        x += float(w["sweep_step"])
    frames = []
    for factor in values:
        frame = score(spec, factor)
        frame["wellbys_per_year"] = factor
        frames.append(frame)
    return pd.concat(frames, ignore_index=True)


def crossovers(swept: pd.DataFrame, spec: dict) -> pd.DataFrame:
    """Smallest WELLBYs-per-year at which a life-saving program beats StrongMinds."""
    strong = float(
        next(c["wellbys_per_1000_usd"] for c in spec["charities"] if c["id"] == "strongminds")
    )
    cash = float(
        next(c["wellbys_per_1000_usd"] for c in spec["charities"] if c["id"] == "givedirectly")
    )
    rows = []
    life = swept[swept["kind"] == "life_saving"]
    for charity_id, group in life.groupby("id"):
        per_year_at_one = float(group["life_years_per_1000"].iloc[0])
        row = group.iloc[0]
        rows.append(
            {
                "id": charity_id,
                "charity": row["charity"],
                "geography": row["geography"],
                "life_years_per_1000": per_year_at_one,
                "wellbys_per_year_to_pass_strongminds": strong / per_year_at_one,
                "wellbys_per_year_to_pass_givedirectly": cash / per_year_at_one,
            }
        )
    return pd.DataFrame(rows).sort_values("life_years_per_1000", ascending=False)


def _label(row: pd.Series) -> str:
    short = {
        "hki_niger": "HKI, Niger",
        "smc_sokoto": "SMC, Sokoto",
        "amf_guinea": "AMF, Guinea",
        "givedirectly": "GiveDirectly",
        "strongminds": "StrongMinds",
    }
    return short[row["id"]]


def plot(baseline: pd.DataFrame, swept: pd.DataFrame, cross: pd.DataFrame, factor: float) -> None:
    FIG.mkdir(parents=True, exist_ok=True)
    order = baseline.sort_values("life_years_per_1000", ascending=True)
    labels = [_label(r) for _, r in order.iterrows()]

    fig, axes = plt.subplots(1, 2, figsize=(9.2, 4.2))
    axes[0].barh(labels, order["life_years_per_1000"], color="#1f4e79")
    axes[0].set_xlabel("Years of life per $1,000")
    axes[0].set_title("Deaths converted to years of life")
    well = baseline.set_index("id").loc[order["id"]]
    axes[1].barh(labels, well["wellbys_per_1000"].to_numpy(), color="#0f6e56")
    axes[1].set_xlabel("WELLBYs per $1,000")
    axes[1].set_title(f"At {factor:g} WELLBYs per year of life")
    fig.tight_layout()
    fig.savefig(FIG / "two_rankings.png", dpi=140)
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(8.2, 4.6))
    colors = {
        "hki_niger": "#1f4e79",
        "smc_sokoto": "#3d7a5a",
        "amf_guinea": "#b45309",
        "givedirectly": "#6b7280",
        "strongminds": "#7c3aed",
    }
    for charity_id, group in swept.groupby("id"):
        ax.plot(
            group["wellbys_per_year"],
            group["wellbys_per_1000"],
            color=colors[charity_id],
            label=_label(group.iloc[0]),
            lw=2,
        )
    ax.axvline(factor, color="#111827", ls=":", lw=1, label=f"Baseline ({factor:g})")
    ax.set_xlabel("WELLBYs credited per year of life saved")
    ax.set_ylabel("WELLBYs per $1,000")
    ax.set_title("Where psychotherapy passes a life-saving program")
    ax.legend(frameon=False, fontsize=8)
    fig.tight_layout()
    fig.savefig(FIG / "wellby_sensitivity.png", dpi=140)
    plt.close(fig)

    cross.to_csv(FIG / "crossovers.csv", index=False)


def main() -> None:
    spec = load_conversions()
    factor = float(spec["wellby"]["wellbys_per_year_of_life"])
    baseline = score(spec, factor, "best")
    median = score(spec, factor, "median")
    swept = sweep(spec)
    cross = crossovers(swept, spec)
    cross_median = crossovers(median, spec)
    baseline.to_csv(HERE / "ranking.csv", index=False)
    median.to_csv(HERE / "ranking_median_geography.csv", index=False)
    plot(baseline, swept, cross, factor)
    summary = {
        "wellbys_per_year_of_life": factor,
        "years_per_under5_malaria_death": spec["life_years"]["years_per_under5_malaria_death"],
        "ranking_best_geography": baseline.to_dict(orient="records"),
        "ranking_median_geography": median.to_dict(orient="records"),
        "crossovers_best_geography": cross.to_dict(orient="records"),
        "crossovers_median_geography": cross_median.to_dict(orient="records"),
    }
    (FIG / "ranking_summary.json").write_text(json.dumps(summary, indent=2))
    print("BEST")
    print(baseline.to_string(index=False))
    print("MEDIAN")
    print(median.to_string(index=False))
    print(cross.to_string(index=False))
    print(cross_median.to_string(index=False))


if __name__ == "__main__":
    main()
