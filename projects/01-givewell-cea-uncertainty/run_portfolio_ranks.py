"""Portfolio rank-order under uncertainty for GiveWell top-charity geographies.

Reads local copies of GiveWell workbooks in data/raw/givewell/, writes
portfolio_ce_ranges.csv and figures/portfolio_*.png.
"""

from __future__ import annotations

import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.stats import spearmanr

ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "data" / "raw" / "givewell"
OUT = Path(__file__).resolve().parent
FIG = OUT / "figures"
N_DRAWS = 20_000
SEED = 42

# GiveWell funding bars (× benchmark). As of May 2026 the published bar is 6×.
# 8× is the prior Top Charities–oriented bar (useful robustness).
# This is a CE screen proxy only — the Top Charities Fund also uses judgment,
# room for more funding, urgency, and implementation quality.
TCF_BAR_PRIMARY = 6.0
TCF_BAR_LEGACY = 8.0
FUNDING_BARS = (TCF_BAR_PRIMARY, TCF_BAR_LEGACY)


def load_moral_weights_points(path: Path) -> pd.DataFrame:
    wb = load_workbook(path, data_only=True)
    ws = wb["Your moral weights"]
    rows = []
    for r in range(6, 200):
        charity = ws.cell(r, 13).value
        if not charity:
            continue
        country = ws.cell(r, 14).value
        region = ws.cell(r, 15).value
        ce = ws.cell(r, 16).value
        if not isinstance(ce, (int, float)):
            continue
        reg = None
        if region is not None and str(region).strip() and str(region).strip().lower() != "nan":
            reg = str(region).strip().replace("FCT (Abuja)", "FCT")
        rows.append(
            {
                "charity": str(charity).strip(),
                "country": str(country).strip() if country else "",
                "region": reg,
                "ce_best": float(ce),
            }
        )
    wb.close()
    return pd.DataFrame(rows)


def extract_sensitivity_cis(path: Path, charity: str) -> pd.DataFrame:
    wb = load_workbook(path, data_only=True)
    if "Sensitivity analysis" not in wb.sheetnames:
        wb.close()
        return pd.DataFrame()
    ws = wb["Sensitivity analysis"]
    starts = [c for c in range(1, 100) if ws.cell(2, c).value]
    target_row = None
    for r in range(1, 60):
        if "Final cost-effectiveness estimate, 25th" in str(ws.cell(r, 1).value or ""):
            target_row = r + 1
            break
    if target_row is None:
        for r in range(20, 60):
            if str(ws.cell(r, 5).value or "") == "Final cost-effectiveness estimate":
                target_row = r
                break
    rows = []
    if target_row is None:
        wb.close()
        return pd.DataFrame()
    for start in starts:
        country = str(ws.cell(2, start).value).strip()
        region = ws.cell(3, start).value
        if isinstance(region, str):
            region = region.strip()
            if region in ("Simple CEA",) or region.lower().startswith("simple"):
                region = None
            else:
                region = region.replace("FCT (Abuja)", "FCT")
        else:
            region = None
        ce_best = ws.cell(5, start).value
        v25 = ws.cell(target_row, start).value
        v75 = ws.cell(target_row, start + 2).value
        if not isinstance(ce_best, (int, float)):
            continue
        if not isinstance(v25, (int, float)) or not isinstance(v75, (int, float)):
            continue
        rows.append(
            {
                "charity": charity,
                "country": country,
                "region": region,
                "ce_best_cea": float(ce_best),
                "ce_25": float(v25),
                "ce_75": float(v75),
                "ci_source": "sensitivity_joint",
            }
        )
    wb.close()
    return pd.DataFrame(rows)


def build_portfolio() -> pd.DataFrame:
    portfolio = load_moral_weights_points(RAW / "moral-weights-tool.xlsx")
    cis = pd.concat(
        [
            extract_sensitivity_cis(RAW / "amf-cea.xlsx", "Against Malaria Foundation"),
            extract_sensitivity_cis(RAW / "hki-vas-cea.xlsx", "Helen Keller International"),
            extract_sensitivity_cis(
                RAW / "malaria-consortium-cea.xlsx", "Malaria Consortium"
            ),
        ],
        ignore_index=True,
    )
    merged = portfolio.merge(cis, on=["charity", "country", "region"], how="left")
    matched = merged["ce_25"].notna()
    base = merged["ce_best_cea"].fillna(merged["ce_best"])
    merged.loc[matched, "rel_lo"] = merged.loc[matched, "ce_25"] / base[matched]
    merged.loc[matched, "rel_hi"] = merged.loc[matched, "ce_75"] / base[matched]
    global_lo = float(merged.loc[matched, "rel_lo"].median())
    global_hi = float(merged.loc[matched, "rel_hi"].median())
    charity_lo = merged.loc[matched].groupby("charity")["rel_lo"].median().to_dict()
    charity_hi = merged.loc[matched].groupby("charity")["rel_hi"].median().to_dict()

    for i, r in merged.iterrows():
        if pd.isna(r["ce_25"]):
            lo = charity_lo.get(r["charity"], global_lo)
            hi = charity_hi.get(r["charity"], global_hi)
            if pd.isna(lo):
                lo, hi = global_lo, global_hi
            merged.at[i, "ce_25"] = r["ce_best"] * float(lo)
            merged.at[i, "ce_75"] = r["ce_best"] * float(hi)
            merged.at[i, "ci_source"] = "imputed_rel_width"
        elif pd.notna(r["ce_best_cea"]) and r["ce_best_cea"] > 0:
            scale = r["ce_best"] / r["ce_best_cea"]
            merged.at[i, "ce_25"] = r["ce_25"] * scale
            merged.at[i, "ce_75"] = r["ce_75"] * scale

    merged["ce_25"] = np.minimum(merged["ce_25"].astype(float), merged["ce_best"] * 0.999)
    merged["ce_75"] = np.maximum(merged["ce_75"].astype(float), merged["ce_best"] * 1.001)
    merged["initiative"] = merged.apply(
        lambda r: f"{r['charity']} — {r['country']}"
        + (f" / {r['region']}" if pd.notna(r["region"]) and r["region"] else ""),
        axis=1,
    )
    return merged.sort_values("ce_best", ascending=False).reset_index(drop=True)


def _ranks_among_eligible(samples: np.ndarray, eligible: np.ndarray) -> np.ndarray:
    """Per-draw rank among geos that clear the bar (1 = best among eligible).

    Ineligible geos get NaN for that draw.
    """
    n, k = samples.shape
    out = np.full((n, k), np.nan)
    for t in range(n):
        mask = eligible[t]
        if not mask.any():
            continue
        idxs = np.where(mask)[0]
        order = idxs[np.argsort(-samples[t, idxs])]
        out[t, order] = np.arange(1, len(order) + 1)
    return out


def run_rank_mc(df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, dict, np.ndarray]:
    df = df.copy()
    df["rank_point"] = np.arange(1, len(df) + 1)
    rng = np.random.default_rng(SEED)
    k = len(df)
    samples = np.column_stack(
        [
            rng.triangular(df.at[i, "ce_25"], df.at[i, "ce_best"], df.at[i, "ce_75"], size=N_DRAWS)
            for i in range(k)
        ]
    )
    order = np.argsort(-samples, axis=1)
    ranks = np.empty_like(order)
    for t in range(N_DRAWS):
        ranks[t, order[t]] = np.arange(1, k + 1)

    df["mean_rank_mc"] = ranks.mean(axis=0)
    df["median_rank_mc"] = np.median(ranks, axis=0)
    df["p_rank1"] = (ranks == 1).mean(axis=0)
    df["p_top5"] = (ranks <= 5).mean(axis=0)
    df["p_top10"] = (ranks <= 10).mean(axis=0)

    # --- Top Charities Fund CE-bar screen (proxy) ---
    for bar in FUNDING_BARS:
        tag = f"{bar:g}x"
        eligible = samples >= bar
        df[f"above_bar_{tag}_point"] = df["ce_best"] >= bar
        df[f"p_clear_bar_{tag}"] = eligible.mean(axis=0)
        ranks_elig = _ranks_among_eligible(samples, eligible)
        with np.errstate(all="ignore"):
            df[f"median_rank_among_eligible_{tag}"] = np.nanmedian(ranks_elig, axis=0)
            df[f"mean_rank_among_eligible_{tag}"] = np.nanmean(ranks_elig, axis=0)
            df[f"p_top5_among_eligible_{tag}"] = np.nanmean(ranks_elig <= 5, axis=0)

    # Rank initiatives by probability of clearing the primary bar (stable screen)
    df["rank_by_p_clear_6x"] = (
        df["p_clear_bar_6x"].rank(ascending=False, method="min").astype(int)
    )

    charities = sorted(df["charity"].unique())
    char_idx = {ch: df.index[df["charity"] == ch].tolist() for ch in charities}
    char_samples = np.column_stack(
        [samples[:, idxs].max(axis=1) for idxs in char_idx.values()]
    )
    char_order = np.argsort(-char_samples, axis=1)
    char_ranks = np.empty_like(char_order)
    for t in range(N_DRAWS):
        char_ranks[t, char_order[t]] = np.arange(1, len(charities) + 1)

    eligible_6 = samples >= TCF_BAR_PRIMARY
    eligible_8 = samples >= TCF_BAR_LEGACY
    char_rows = []
    for i, ch in enumerate(charities):
        idxs = char_idx[ch]
        char_rows.append(
            {
                "charity": ch,
                "point_best_geo_ce": df.loc[idxs, "ce_best"].max(),
                "n_geos": len(idxs),
                "n_above_6x_point": int((df.loc[idxs, "ce_best"] >= TCF_BAR_PRIMARY).sum()),
                "n_above_8x_point": int((df.loc[idxs, "ce_best"] >= TCF_BAR_LEGACY).sum()),
                "p_rank1": float((char_ranks[:, i] == 1).mean()),
                "median_rank_mc": float(np.median(char_ranks[:, i])),
                "mean_geos_clearing_6x": float(eligible_6[:, idxs].sum(axis=1).mean()),
                "mean_geos_clearing_8x": float(eligible_8[:, idxs].sum(axis=1).mean()),
                "p_any_geo_clears_6x": float(eligible_6[:, idxs].any(axis=1).mean()),
                "p_any_geo_clears_8x": float(eligible_8[:, idxs].any(axis=1).mean()),
            }
        )
    char_summary = pd.DataFrame(char_rows).sort_values(
        "point_best_geo_ce", ascending=False
    )

    rho, _ = spearmanr(df["rank_point"], df["median_rank_mc"])
    point_top5 = set(df.nsmallest(5, "rank_point")["initiative"])
    jaccard = float(
        np.mean(
            [
                len(point_top5 & set(df.iloc[order[t, :5]]["initiative"]))
                / len(point_top5 | set(df.iloc[order[t, :5]]["initiative"]))
                for t in range(N_DRAWS)
            ]
        )
    )
    n_clear_6 = eligible_6.sum(axis=1)
    n_clear_8 = eligible_8.sum(axis=1)
    summary = {
        "n_initiatives": int(k),
        "n_draws": N_DRAWS,
        "metric": "x_givewell_benchmark",
        "tcf_bar_primary": TCF_BAR_PRIMARY,
        "tcf_bar_legacy": TCF_BAR_LEGACY,
        "tcf_bar_note": (
            "Proxy for Top Charities Fund CE screen only. GiveWell also uses "
            "judgment, room for more funding, and implementation factors. "
            "All 87 geos are already Top Charity programs; the screen is "
            "whether the geography clears the ×-benchmark bar under uncertainty."
        ),
        "point_estimate_leader": df.at[0, "initiative"],
        "p_leader_remains_rank1": float(df.at[0, "p_rank1"]),
        "spearman_point_rank_vs_median_mc_rank": float(rho),
        "mean_jaccard_top5_vs_point_top5": jaccard,
        "native_ci_count": int((df["ci_source"] == "sensitivity_joint").sum()),
        "imputed_ci_count": int((df["ci_source"] == "imputed_rel_width").sum()),
        "charity_level_p_rank1": {
            r.charity: round(float(r.p_rank1), 4) for r in char_summary.itertuples()
        },
        "n_above_6x_point": int((df["ce_best"] >= TCF_BAR_PRIMARY).sum()),
        "n_above_8x_point": int((df["ce_best"] >= TCF_BAR_LEGACY).sum()),
        "mean_n_clearing_6x": float(n_clear_6.mean()),
        "p5_n_clearing_6x": float(np.quantile(n_clear_6, 0.05)),
        "p95_n_clearing_6x": float(np.quantile(n_clear_6, 0.95)),
        "mean_n_clearing_8x": float(n_clear_8.mean()),
        "p5_n_clearing_8x": float(np.quantile(n_clear_8, 0.05)),
        "p95_n_clearing_8x": float(np.quantile(n_clear_8, 0.95)),
        "n_always_clear_6x": int((df["p_clear_bar_6x"] >= 0.99).sum()),
        "n_rarely_clear_6x": int((df["p_clear_bar_6x"] < 0.10).sum()),
        "n_always_clear_8x": int((df["p_clear_bar_8x"] >= 0.99).sum()),
        "n_rarely_clear_8x": int((df["p_clear_bar_8x"] < 0.10).sum()),
    }
    return df, char_summary, summary, samples


def save_figures(
    df: pd.DataFrame,
    char_summary: pd.DataFrame,
    rho: float,
    samples: np.ndarray,
) -> None:
    FIG.mkdir(parents=True, exist_ok=True)
    k = len(df)

    plot_df = df.head(20).iloc[::-1]
    fig, ax = plt.subplots(figsize=(9.5, 7.5))
    y = np.arange(len(plot_df))
    ax.hlines(y, plot_df["ce_25"], plot_df["ce_75"], color="#6b7c8c", linewidth=2)
    ax.plot(
        plot_df["ce_best"],
        y,
        "o",
        color="#c47b3a",
        markersize=6,
        label="GiveWell point estimate",
    )
    ax.axvline(TCF_BAR_PRIMARY, color="#2f5d4a", linestyle="--", linewidth=1.2, label="6× funding bar")
    ax.axvline(TCF_BAR_LEGACY, color="#8a6d3b", linestyle=":", linewidth=1.2, label="8× legacy bar")
    ax.set_yticks(y)
    ax.set_yticklabels(plot_df["initiative"], fontsize=8)
    ax.set_xlabel("Cost-effectiveness (× GiveWell benchmark)")
    ax.set_title("Top 20 program geographies: CE and uncertainty range")
    ax.legend(frameon=False, loc="lower right")
    fig.tight_layout()
    fig.savefig(FIG / "portfolio_ce_ranges.png", dpi=160)
    plt.close(fig)

    p1 = df.nlargest(12, "p_rank1").iloc[::-1]
    fig, ax = plt.subplots(figsize=(8, 5))
    ax.barh(p1["initiative"], p1["p_rank1"], color="#3b6d9a")
    ax.set_xlabel("Share of draws ranked #1")
    ax.set_title("Probability of being the top initiative under uncertainty")
    fig.tight_layout()
    fig.savefig(FIG / "portfolio_p_rank1.png", dpi=160)
    plt.close(fig)

    colors = {
        "Against Malaria Foundation": "#3b6d9a",
        "Helen Keller International": "#c47b3a",
        "Malaria Consortium": "#5a8f6b",
        "New Incentives": "#8b6b9a",
    }
    fig, ax = plt.subplots(figsize=(6, 6))
    for ch, g in df.groupby("charity"):
        ax.scatter(
            g["rank_point"],
            g["median_rank_mc"],
            c=colors.get(ch, "#333"),
            alpha=0.75,
            label=ch,
            edgecolors="none",
            s=28,
        )
    ax.plot([1, k], [1, k], "--", color="gray", linewidth=1)
    ax.set_xlabel("Rank by point estimate (1 = best)")
    ax.set_ylabel("Median rank under Monte Carlo")
    ax.set_title(f"Rank-order stability (ρ = {rho:.2f})")
    ax.legend(fontsize=7, frameon=False)
    ax.set_aspect("equal")
    fig.tight_layout()
    fig.savefig(FIG / "portfolio_rank_scatter.png", dpi=160)
    plt.close(fig)

    cs = char_summary.sort_values("p_rank1")
    fig, ax = plt.subplots(figsize=(7, 3.8))
    ax.barh(cs["charity"], cs["p_rank1"], color="#3b6d9a")
    ax.set_xlabel("P(charity owns the single highest-CE geography)")
    ax.set_title("Charity-level leadership under uncertainty")
    fig.tight_layout()
    fig.savefig(FIG / "portfolio_charity_p_rank1.png", dpi=160)
    plt.close(fig)

    # P(clear 6× bar) — all geos, sorted
    clear = df.sort_values("p_clear_bar_6x").copy()
    fig, ax = plt.subplots(figsize=(9, 11))
    bar_colors = [colors.get(c, "#333") for c in clear["charity"]]
    ax.barh(clear["initiative"], clear["p_clear_bar_6x"], color=bar_colors)
    ax.axvline(0.5, color="gray", linestyle="--", linewidth=0.8)
    ax.set_xlabel("P(CE ≥ 6× benchmark) under Monte Carlo")
    ax.set_title("Top Charities Fund CE-screen proxy: probability of clearing the 6× bar")
    ax.set_xlim(0, 1.02)
    ax.tick_params(axis="y", labelsize=6)
    fig.tight_layout()
    fig.savefig(FIG / "portfolio_p_clear_6x.png", dpi=160)
    plt.close(fig)

    # Marginal / contested geos: not always clear, not never clear
    contested = df[(df["p_clear_bar_6x"] > 0.05) & (df["p_clear_bar_6x"] < 0.95)].sort_values(
        "p_clear_bar_6x"
    )
    if len(contested) > 0:
        fig, ax = plt.subplots(figsize=(8, max(3.5, 0.28 * len(contested))))
        ax.barh(
            contested["initiative"],
            contested["p_clear_bar_6x"],
            color=[colors.get(c, "#333") for c in contested["charity"]],
        )
        ax.axvline(0.5, color="gray", linestyle="--", linewidth=0.8)
        ax.set_xlabel("P(CE ≥ 6×)")
        ax.set_title("Contested geographies (5% < P(clear 6×) < 95%)")
        ax.set_xlim(0, 1.02)
        ax.tick_params(axis="y", labelsize=7)
        fig.tight_layout()
        fig.savefig(FIG / "portfolio_p_clear_6x_contested.png", dpi=160)
        plt.close(fig)

    # How many geos clear the bar each draw
    n6 = (samples >= TCF_BAR_PRIMARY).sum(axis=1)
    n8 = (samples >= TCF_BAR_LEGACY).sum(axis=1)
    fig, ax = plt.subplots(figsize=(8, 4.2))
    ax.hist(n6, bins=range(int(n6.min()), int(n6.max()) + 2), color="#3b6d9a", alpha=0.85, label="≥ 6×")
    ax.hist(n8, bins=range(int(n8.min()), int(n8.max()) + 2), color="#c47b3a", alpha=0.55, label="≥ 8×")
    ax.axvline(n6.mean(), color="#3b6d9a", linestyle="--", linewidth=1.2)
    ax.axvline(n8.mean(), color="#c47b3a", linestyle="--", linewidth=1.2)
    ax.set_xlabel("Number of initiative-geographies clearing the bar")
    ax.set_ylabel("Monte Carlo draws")
    ax.set_title("Size of the CE-eligible set under uncertainty")
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(FIG / "portfolio_n_clearing_bar.png", dpi=160)
    plt.close(fig)

    # Charity mean geos clearing 6×
    cs2 = char_summary.sort_values("mean_geos_clearing_6x")
    fig, ax = plt.subplots(figsize=(7, 3.8))
    ax.barh(cs2["charity"], cs2["mean_geos_clearing_6x"], color="#3b6d9a")
    ax.set_xlabel("Mean # of that charity’s geos with CE ≥ 6× (per draw)")
    ax.set_title("Which Top Charities populate the CE-eligible set?")
    fig.tight_layout()
    fig.savefig(FIG / "portfolio_charity_mean_clear_6x.png", dpi=160)
    plt.close(fig)


def main() -> None:
    FIG.mkdir(parents=True, exist_ok=True)
    df = build_portfolio()
    df, char_summary, summary, samples = run_rank_mc(df)
    cols = [
        "rank_point",
        "rank_by_p_clear_6x",
        "initiative",
        "charity",
        "country",
        "region",
        "ce_best",
        "ce_25",
        "ce_75",
        "above_bar_6x_point",
        "above_bar_8x_point",
        "p_clear_bar_6x",
        "p_clear_bar_8x",
        "median_rank_among_eligible_6x",
        "p_top5_among_eligible_6x",
        "median_rank_among_eligible_8x",
        "p_top5_among_eligible_8x",
        "p_rank1",
        "p_top5",
        "p_top10",
        "median_rank_mc",
        "mean_rank_mc",
        "ci_source",
    ]
    df[cols].to_csv(OUT / "portfolio_ce_ranges.csv", index=False)
    char_summary.to_csv(OUT / "portfolio_charity_level.csv", index=False)
    # Compact fund-eligibility table for the brief
    fund_cols = [
        "rank_by_p_clear_6x",
        "initiative",
        "charity",
        "ce_best",
        "p_clear_bar_6x",
        "p_clear_bar_8x",
        "median_rank_among_eligible_6x",
        "above_bar_6x_point",
        "above_bar_8x_point",
        "ci_source",
    ]
    df.sort_values("rank_by_p_clear_6x")[fund_cols].to_csv(
        OUT / "portfolio_fund_eligibility.csv", index=False
    )
    (FIG / "portfolio_rank_summary.json").write_text(json.dumps(summary, indent=2))
    save_figures(df, char_summary, summary["spearman_point_rank_vs_median_mc_rank"], samples)
    print(json.dumps(summary, indent=2))
    print("\nLowest P(clear 6×) among point-above-6×:")
    weak = df[df["above_bar_6x_point"]].nsmallest(8, "p_clear_bar_6x")
    print(weak[["initiative", "ce_best", "p_clear_bar_6x", "p_clear_bar_8x"]].to_string(index=False))
    print("\nHighest P(clear 6×) among point-below-6×:")
    climb = df[~df["above_bar_6x_point"]].nlargest(8, "p_clear_bar_6x")
    print(climb[["initiative", "ce_best", "p_clear_bar_6x"]].to_string(index=False))


if __name__ == "__main__":
    main()
